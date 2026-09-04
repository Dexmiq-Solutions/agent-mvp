"""Comprehensive unit tests for the Parsing & Extraction layer."""

import io
import docx
import pytest

from acquisition.formats import DocumentType
from exceptions.parsing import (
    DocumentExtractionError,
    InvalidParsingInputError,
    ParsingError,
    UnsupportedDocumentTypeError,
)
from ingestion.models import IngestedDocument
from parsing import (
    BaseParser,
    DOCXParser,
    DocumentParsingService,
    ElementType,
    MarkdownParser,
    ParsedDocument,
    ParsedElement,
    ParserRegistry,
    TXTParser,
    get_parser_registry,
    get_parsing_service,
    reset_parser_registry,
    reset_parsing_service,
)


@pytest.fixture(autouse=True)
def cleanup_singletons():
    """Reset singletons before and after each test."""
    reset_parser_registry()
    reset_parsing_service()
    yield
    reset_parser_registry()
    reset_parsing_service()


def make_ingested_doc(
    content: bytes | str,
    doc_type: DocumentType,
    doc_id: str = "doc-test-1",
    project_id: str = "proj-test",
    filename: str = "testfile",
    content_type: str = "text/plain",
) -> IngestedDocument:
    """Helper to construct an IngestedDocument test fixture."""
    raw_bytes = content.encode("utf-8") if isinstance(content, str) else content
    return IngestedDocument(
        document_id=doc_id,
        project_id=project_id,
        source_storage_path=f"{project_id}/{filename}",
        original_filename=filename,
        detected_document_type=doc_type,
        content_type=content_type,
        raw_bytes=raw_bytes,
        source_metadata={"test_source": True},
        document_version_id="v1",
        size_bytes=len(raw_bytes),
    )


# ==============================================================================
# 1. TXT Parser Tests
# ==============================================================================


def test_txt_parser_normal_paragraphs():
    """Verify TXT parser extracts ordered text blocks separated by newlines."""
    parser = TXTParser()
    text = (
        "The system uses OAuth for authentication.\n\n"
        "Access tokens expire after 60 minutes.\n\n"
        "Refresh tokens are valid for 30 days."
    )
    doc = make_ingested_doc(text, DocumentType.TXT, filename="auth.txt")
    result = parser.parse(doc)

    assert isinstance(result, ParsedDocument)
    assert result.document_type == DocumentType.TXT
    assert result.document_id == "doc-test-1"
    assert result.project_id == "proj-test"
    assert result.total_elements == 3

    assert result.elements[0].content == "The system uses OAuth for authentication."
    assert result.elements[0].order == 0
    assert result.elements[0].element_type == ElementType.TEXT_BLOCK
    assert result.elements[0].heading_level is None
    assert result.elements[0].parent_id is None
    assert result.elements[0].section_path == ()

    assert result.elements[1].content == "Access tokens expire after 60 minutes."
    assert result.elements[1].order == 1

    assert result.elements[2].content == "Refresh tokens are valid for 30 days."
    assert result.elements[2].order == 2


def test_txt_parser_minimal_document():
    """Verify TXT parser handles a single-line document cleanly."""
    parser = TXTParser()
    doc = make_ingested_doc("Single line content without newlines.", DocumentType.TXT)
    result = parser.parse(doc)

    assert result.total_elements == 1
    assert result.elements[0].content == "Single line content without newlines."
    assert result.elements[0].order == 0


def test_txt_parser_empty_content():
    """Verify TXT parser returns zero elements for empty or whitespace-only documents."""
    parser = TXTParser()
    doc_empty = make_ingested_doc("", DocumentType.TXT)
    res_empty = parser.parse(doc_empty)
    assert res_empty.total_elements == 0
    assert res_empty.elements == []

    doc_ws = make_ingested_doc("   \n\n   \t\n  ", DocumentType.TXT)
    res_ws = parser.parse(doc_ws)
    assert res_ws.total_elements == 0


def test_txt_parser_decoding_fallbacks():
    """Verify TXT parser decodes UTF-8-BOM and Latin-1 correctly without failing."""
    parser = TXTParser()

    # UTF-8 with BOM
    bom_content = "\ufeffHello with UTF-8 BOM\n\nSecond block.".encode("utf-8-sig")
    doc_bom = make_ingested_doc(bom_content, DocumentType.TXT)
    res_bom = parser.parse(doc_bom)
    assert res_bom.total_elements == 2
    assert "Hello with UTF-8 BOM" in res_bom.elements[0].content

    # Latin-1 encoded bytes (contains accents)
    latin1_bytes = "Café au lait\n\nTrès bien".encode("latin-1")
    doc_latin1 = make_ingested_doc(latin1_bytes, DocumentType.TXT)
    res_latin1 = parser.parse(doc_latin1)
    assert res_latin1.total_elements == 2
    assert "Café" in res_latin1.elements[0].content


def test_txt_parser_convenience_parse_content():
    """Verify parse_content helper works directly with strings."""
    parser = TXTParser()
    result = parser.parse_content("Block one\n\nBlock two")
    assert result.total_elements == 2
    assert result.document_type == DocumentType.TXT


# ==============================================================================
# 2. Markdown Parser Tests
# ==============================================================================


def test_markdown_parser_nested_hierarchy():
    """Verify Markdown parser preserves H1 -> H2 -> H3 nested hierarchies and parent pointers."""
    parser = MarkdownParser()
    md_text = (
        "# Authentication\n\n"
        "Users authenticate using OAuth.\n\n"
        "## Token Expiration\n\n"
        "Access tokens expire after 60 minutes.\n\n"
        "### Grace Period\n\n"
        "There is a 5 minute grace period for clock skew.\n\n"
        "## Multi-Factor\n\n"
        "MFA is enforced for admin users."
    )
    doc = make_ingested_doc(md_text, DocumentType.MARKDOWN, filename="security.md")
    result = parser.parse(doc)

    assert result.total_elements == 8
    elements = result.elements

    # 1. H1: Authentication
    h1 = elements[0]
    assert h1.element_type == ElementType.HEADING
    assert h1.content == "Authentication"
    assert h1.heading_level == 1
    assert h1.parent_id is None
    assert h1.section_path == ("Authentication",)

    # 2. Paragraph under H1
    p1 = elements[1]
    assert p1.element_type == ElementType.PARAGRAPH
    assert p1.content == "Users authenticate using OAuth."
    assert p1.parent_id == h1.element_id
    assert p1.section_path == ("Authentication",)

    # 3. H2: Token Expiration (child of H1)
    h2_1 = elements[2]
    assert h2_1.element_type == ElementType.HEADING
    assert h2_1.content == "Token Expiration"
    assert h2_1.heading_level == 2
    assert h2_1.parent_id == h1.element_id
    assert h2_1.section_path == ("Authentication", "Token Expiration")

    # 4. Paragraph under H2
    p2 = elements[3]
    assert p2.element_type == ElementType.PARAGRAPH
    assert p2.content == "Access tokens expire after 60 minutes."
    assert p2.parent_id == h2_1.element_id
    assert p2.section_path == ("Authentication", "Token Expiration")

    # 5. H3: Grace Period (child of H2)
    h3 = elements[4]
    assert h3.element_type == ElementType.HEADING
    assert h3.content == "Grace Period"
    assert h3.heading_level == 3
    assert h3.parent_id == h2_1.element_id
    assert h3.section_path == ("Authentication", "Token Expiration", "Grace Period")

    # 6. Paragraph under H3
    p3 = elements[5]
    assert p3.element_type == ElementType.PARAGRAPH
    assert p3.content == "There is a 5 minute grace period for clock skew."
    assert p3.parent_id == h3.element_id
    assert p3.section_path == ("Authentication", "Token Expiration", "Grace Period")

    # 7. H2: Multi-Factor (sibling of Token Expiration, child of H1)
    h2_2 = elements[6]
    assert h2_2.element_type == ElementType.HEADING
    assert h2_2.content == "Multi-Factor"
    assert h2_2.heading_level == 2
    assert h2_2.parent_id == h1.element_id
    assert h2_2.section_path == ("Authentication", "Multi-Factor")

    # 8. Paragraph under H2 Multi-Factor
    p4 = elements[7]
    assert p4.element_type == ElementType.PARAGRAPH
    assert p4.content == "MFA is enforced for admin users."
    assert p4.parent_id == h2_2.element_id
    assert p4.section_path == ("Authentication", "Multi-Factor")


def test_markdown_parser_level_skipping():
    """Verify Markdown parser handles level skipping (e.g. H1 followed directly by H3)."""
    parser = MarkdownParser()
    md_text = (
        "# Top Level\n\n"
        "Top text.\n\n"
        "### Deep Subsection\n\n"
        "Deep text."
    )
    result = parser.parse_content(md_text)
    assert result.total_elements == 4

    h1 = result.elements[0]
    h3 = result.elements[2]
    p_deep = result.elements[3]

    assert h3.heading_level == 3
    assert h3.parent_id == h1.element_id
    assert h3.section_path == ("Top Level", "Deep Subsection")
    assert p_deep.parent_id == h3.element_id


def test_markdown_code_blocks_ignore_headings():
    """Verify # inside fenced code blocks is not parsed as a heading."""
    parser = MarkdownParser()
    md_text = (
        "# Python Example\n\n"
        "Here is some code:\n\n"
        "```python\n"
        "# This is a comment, NOT a heading\n"
        "def hello():\n"
        "    ## Another comment\n"
        "    return True\n"
        "```\n\n"
        "Paragraph after code block."
    )
    result = parser.parse_content(md_text)

    # Should have: H1, Paragraph, Code Block, Paragraph
    assert result.total_elements == 4
    assert len(result.headings) == 1
    assert result.headings[0].content == "Python Example"

    code_elem = result.elements[2]
    assert code_elem.element_type == ElementType.CODE_BLOCK
    assert "# This is a comment, NOT a heading" in code_elem.content
    assert code_elem.metadata.get("language") == "python"
    assert code_elem.section_path == ("Python Example",)


def test_markdown_table_detection():
    """Verify Markdown tables are detected as ElementType.TABLE."""
    parser = MarkdownParser()
    md_text = (
        "# Data Summary\n\n"
        "| ID | Status |\n"
        "|---|---|\n"
        "| 1 | Active |\n"
        "| 2 | Inactive |"
    )
    result = parser.parse_content(md_text)
    assert result.total_elements == 2
    assert result.elements[0].element_type == ElementType.HEADING
    assert result.elements[1].element_type == ElementType.TABLE
    assert "| 1 | Active |" in result.elements[1].content


def test_markdown_parser_empty_content():
    """Verify Markdown parser handles empty or whitespace documents without error."""
    parser = MarkdownParser()
    res = parser.parse_content("   \n\n   ")
    assert res.total_elements == 0
    assert res.elements == []


# ==============================================================================
# 3. DOCX Parser Tests
# ==============================================================================


def create_sample_docx_bytes() -> bytes:
    """Generate an in-memory DOCX document with headings, paragraphs, and a table."""
    doc = docx.Document()
    doc.add_heading("Authentication", level=1)
    doc.add_paragraph("Users authenticate using OAuth.")
    doc.add_heading("Token Expiration", level=2)
    doc.add_paragraph("Tokens expire after 60 minutes.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Timeout"
    table.cell(1, 0).text = "Admin"
    table.cell(1, 1).text = "15m"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_parser_normal_document():
    """Verify DOCX parser extracts Word headings, paragraphs, and tables preserving hierarchy."""
    parser = DOCXParser()
    raw_bytes = create_sample_docx_bytes()
    doc = make_ingested_doc(raw_bytes, DocumentType.DOCX, filename="spec.docx")
    result = parser.parse(doc)

    assert isinstance(result, ParsedDocument)
    assert result.document_type == DocumentType.DOCX
    assert result.total_elements == 5

    # 1. Heading 1
    h1 = result.elements[0]
    assert h1.element_type == ElementType.HEADING
    assert h1.content == "Authentication"
    assert h1.heading_level == 1
    assert h1.parent_id is None
    assert h1.section_path == ("Authentication",)

    # 2. Paragraph
    p1 = result.elements[1]
    assert p1.element_type == ElementType.PARAGRAPH
    assert p1.content == "Users authenticate using OAuth."
    assert p1.parent_id == h1.element_id
    assert p1.section_path == ("Authentication",)

    # 3. Heading 2
    h2 = result.elements[2]
    assert h2.element_type == ElementType.HEADING
    assert h2.content == "Token Expiration"
    assert h2.heading_level == 2
    assert h2.parent_id == h1.element_id
    assert h2.section_path == ("Authentication", "Token Expiration")

    # 4. Paragraph
    p2 = result.elements[3]
    assert p2.element_type == ElementType.PARAGRAPH
    assert p2.content == "Tokens expire after 60 minutes."
    assert p2.parent_id == h2.element_id
    assert p2.section_path == ("Authentication", "Token Expiration")

    # 5. Table
    tbl = result.elements[4]
    assert tbl.element_type == ElementType.TABLE
    assert "Role | Timeout" in tbl.content
    assert "Admin | 15m" in tbl.content
    assert tbl.parent_id == h2.element_id
    assert tbl.section_path == ("Authentication", "Token Expiration")


def test_docx_parser_title_style():
    """Verify DOCX parser treats Title style as heading level 1."""
    doc = docx.Document()
    doc.add_paragraph("Project Charter", style="Title")
    doc.add_paragraph("Charter introduction.")

    buffer = io.BytesIO()
    doc.save(buffer)

    parser = DOCXParser()
    result = parser.parse_content(buffer.getvalue(), original_filename="charter.docx")
    assert result.total_elements == 2
    assert result.elements[0].element_type == ElementType.HEADING
    assert result.elements[0].heading_level == 1
    assert result.elements[0].content == "Project Charter"
    assert result.elements[1].parent_id == result.elements[0].element_id


def test_docx_parser_empty_document():
    """Verify DOCX parser handles an empty DOCX file cleanly."""
    doc = docx.Document()
    buffer = io.BytesIO()
    doc.save(buffer)

    parser = DOCXParser()
    result = parser.parse_content(buffer.getvalue(), original_filename="empty.docx")
    assert result.total_elements == 0
    assert result.elements == []


def test_docx_parser_corrupted_payload():
    """Verify DOCX parser raises DocumentExtractionError on corrupt bytes."""
    parser = DOCXParser()
    corrupt_bytes = b"not a valid zip or docx payload"
    doc = make_ingested_doc(corrupt_bytes, DocumentType.DOCX, filename="bad.docx")

    with pytest.raises(DocumentExtractionError) as exc_info:
        parser.parse(doc)
    assert "Failed to open or unpack DOCX file" in str(exc_info.value)


# ==============================================================================
# 4. Parser Registry Tests
# ==============================================================================


def test_parser_registry_resolution():
    """Verify ParserRegistry correctly resolves parsers for TXT, Markdown, and DOCX."""
    registry = get_parser_registry()

    assert isinstance(registry.get_parser(DocumentType.TXT), TXTParser)
    assert isinstance(registry.get_parser(DocumentType.MARKDOWN), MarkdownParser)
    assert isinstance(registry.get_parser(DocumentType.DOCX), DOCXParser)

    # Resolution by filename / extension
    assert isinstance(registry.resolve_parser("notes.txt"), TXTParser)
    assert isinstance(registry.resolve_parser("README.md"), MarkdownParser)
    assert isinstance(registry.resolve_parser("spec.docx"), DOCXParser)


def test_parser_registry_unsupported_pdf():
    """Verify ParserRegistry rejects PDF with UnsupportedDocumentTypeError in this phase."""
    registry = get_parser_registry()

    with pytest.raises(UnsupportedDocumentTypeError) as exc_info:
        registry.get_parser(DocumentType.PDF)
    assert "No parser registered for document type 'pdf'" in str(exc_info.value)

    with pytest.raises(UnsupportedDocumentTypeError):
        registry.resolve_parser("document.pdf")


def test_parser_registry_unknown_format():
    """Verify ParserRegistry rejects unmapped file extensions."""
    registry = get_parser_registry()
    with pytest.raises(UnsupportedDocumentTypeError):
        registry.resolve_parser("archive.tar.gz")


def test_parser_registry_extensibility():
    """Verify registry supports registering new parsers without breaking existing ones."""
    registry = ParserRegistry()

    class DummyPDFParser(BaseParser):
        @property
        def supported_document_type(self) -> DocumentType:
            return DocumentType.PDF

        def parse(self, document: IngestedDocument) -> ParsedDocument:
            return ParsedDocument(
                document_id=document.document_id,
                project_id=document.project_id,
                document_type=DocumentType.PDF,
            )

    registry.register_parser(DocumentType.PDF, DummyPDFParser())
    assert registry.is_supported(DocumentType.PDF)
    assert isinstance(registry.get_parser(DocumentType.PDF), DummyPDFParser)


# ==============================================================================
# 5. Structured Output & Representation Tests
# ==============================================================================


def test_parsed_document_properties_and_methods():
    """Verify ParsedDocument helper methods and properties work as expected."""
    elem1 = ParsedElement(
        element_id="e-0",
        element_type=ElementType.HEADING,
        content="Heading 1",
        order=0,
        heading_level=1,
    )
    elem2 = ParsedElement(
        element_id="e-1",
        element_type=ElementType.PARAGRAPH,
        content="Paragraph content.",
        order=1,
    )
    parsed = ParsedDocument(
        document_id="doc-123",
        project_id="proj-456",
        document_type=DocumentType.MARKDOWN,
        elements=[elem1, elem2],
        source_metadata={"author": "Alice"},
    )

    assert parsed.total_elements == 2
    assert len(parsed.headings) == 1
    assert parsed.headings[0].content == "Heading 1"
    assert len(parsed.get_elements_by_type(ElementType.PARAGRAPH)) == 1
    assert parsed.get_full_text() == "Heading 1\n\nParagraph content."

    # Serialization
    doc_dict = parsed.to_dict(include_elements=True)
    assert doc_dict["document_id"] == "doc-123"
    assert doc_dict["project_id"] == "proj-456"
    assert doc_dict["total_elements"] == 2
    assert len(doc_dict["elements"]) == 2
    assert doc_dict["elements"][0]["content"] == "Heading 1"


def test_parsed_document_repr_masks_content():
    """Verify ParsedDocument.__repr__ does not expose document text or leak payloads."""
    secret_text = "TOP_SECRET_PROPRIETARY_INFORMATION_DO_NOT_LOG"
    elem = ParsedElement(
        element_id="e-0",
        element_type=ElementType.PARAGRAPH,
        content=secret_text,
        order=0,
    )
    doc = ParsedDocument(
        document_id="doc-secret",
        project_id="proj-confidential",
        document_type=DocumentType.TXT,
        elements=[elem],
    )

    rep = repr(doc)
    assert secret_text not in rep
    assert "doc-secret" in rep
    assert "proj-confidential" in rep
    assert "total_elements=1" in rep


# ==============================================================================
# 6. Service Layer & Async Thread Offloading Tests
# ==============================================================================


@pytest.mark.asyncio
async def test_parsing_service_async_execution():
    """Verify DocumentParsingService executes parsing asynchronously via thread offloading."""
    service = get_parsing_service()
    doc = make_ingested_doc(
        "# Fast API\n\nHigh performance web framework.",
        DocumentType.MARKDOWN,
        filename="fastapi.md",
    )

    result = await service.parse(doc)
    assert isinstance(result, ParsedDocument)
    assert result.total_elements == 2
    assert result.headings[0].content == "Fast API"


@pytest.mark.asyncio
async def test_parsing_service_batch_execution():
    """Verify DocumentParsingService batch parsing processes multiple document formats concurrently."""
    service = get_parsing_service()

    doc_txt = make_ingested_doc("Text content 1\n\nText content 2", DocumentType.TXT, doc_id="d1")
    doc_md = make_ingested_doc("# Title\n\nMarkdown content", DocumentType.MARKDOWN, doc_id="d2")
    docx_bytes = create_sample_docx_bytes()
    doc_docx = make_ingested_doc(docx_bytes, DocumentType.DOCX, doc_id="d3")

    results = await service.parse_batch([doc_txt, doc_md, doc_docx])
    assert len(results) == 3
    assert results[0].document_type == DocumentType.TXT
    assert results[1].document_type == DocumentType.MARKDOWN
    assert results[2].document_type == DocumentType.DOCX


def test_parsing_service_invalid_input():
    """Verify DocumentParsingService rejects non-IngestedDocument inputs."""
    service = get_parsing_service()
    with pytest.raises(InvalidParsingInputError):
        service.parse_sync("not an ingested document")  # type: ignore
