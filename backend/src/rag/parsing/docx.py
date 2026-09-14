"""Native DOCX document parser."""

import io
import re
from typing import Any

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from rag.acquisition.formats import DocumentType
from exceptions.parsing import DocumentExtractionError, InvalidParsingInputError
from rag.ingestion.models import IngestedDocument
from rag.parsing.base import BaseParser
from rag.parsing.models import ElementType, ParsedDocument, ParsedElement

HEADING_STYLE_PATTERN = re.compile(r"^heading\s+(\d+)$", re.IGNORECASE)


class DOCXParser(BaseParser):
    """Structure-aware parser for Microsoft Word documents (.docx).

    Uses python-docx to extract native Word headings, paragraphs, and tables,
    preserving:
    - Document order
    - Native heading levels (from Heading 1-9 and Title styles)
    - Nested section hierarchy (parent-child links and section breadcrumbs)
    """

    @property
    def supported_document_type(self) -> DocumentType:
        """Return DocumentType.DOCX."""
        return DocumentType.DOCX

    def _extract_heading_level(self, style_name: str) -> int | None:
        """Determine heading level from Word style name.

        Recognizes standard Word styles like 'Heading 1', 'Heading 2', and 'Title'.
        """
        clean = style_name.strip()
        match = HEADING_STYLE_PATTERN.match(clean)
        if match:
            return int(match.group(1))
        lower = clean.lower()
        if lower == "title":
            return 1
        if lower == "subtitle":
            return 2
        return None

    def _extract_table_content(self, table: Table) -> str:
        """Extract table contents into a readable plain-text tabular format."""
        rows_text: list[str] = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(row_cells))
        return "\n".join(rows_text).strip()

    def parse(self, document: IngestedDocument) -> ParsedDocument:
        """Parse an IngestedDocument containing DOCX bytes.

        Args:
            document: IngestedDocument with DOCX file bytes.

        Returns:
            ParsedDocument containing ordered, structured elements.

        Raises:
            InvalidParsingInputError: If document is None or missing raw_bytes.
            DocumentExtractionError: If file is corrupted or cannot be parsed.
        """
        if document is None or document.raw_bytes is None:
            raise InvalidParsingInputError("Document and raw_bytes must not be None.")

        try:
            doc_stream = io.BytesIO(document.raw_bytes)
            word_doc = docx.Document(doc_stream)
        except Exception as exc:
            raise DocumentExtractionError(
                f"Failed to open or unpack DOCX file '{document.document_id}': {str(exc)}",
                original_error=exc,
            ) from exc

        try:
            elements: list[ParsedElement] = []
            # Stack stores tuples of: (heading_level, element_id, heading_title)
            heading_stack: list[tuple[int, str, str]] = []

            # Retrieve inner content blocks (paragraphs and tables) in exact sequence
            if hasattr(word_doc, "iter_inner_content"):
                content_blocks = list(word_doc.iter_inner_content())
            else:
                content_blocks = list(word_doc.paragraphs)

            for block in content_blocks:
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue

                    style_name = block.style.name if block.style and block.style.name else ""
                    heading_level = self._extract_heading_level(style_name)

                    if heading_level is not None:
                        # Manage hierarchy stack
                        while heading_stack and heading_stack[-1][0] >= heading_level:
                            heading_stack.pop()

                        parent_id = heading_stack[-1][1] if heading_stack else None
                        section_path = tuple(item[2] for item in heading_stack) + (text,)
                        order = len(elements)
                        element_id = f"elem-{order}"

                        element = ParsedElement(
                            element_id=element_id,
                            element_type=ElementType.HEADING,
                            content=text,
                            order=order,
                            heading_level=heading_level,
                            parent_id=parent_id,
                            section_path=section_path,
                            metadata={
                                "style_name": style_name,
                                "heading_level": heading_level,
                                "character_count": len(text),
                            },
                        )
                        elements.append(element)
                        heading_stack.append((heading_level, element_id, text))
                    else:
                        parent_id = heading_stack[-1][1] if heading_stack else None
                        section_path = tuple(item[2] for item in heading_stack)
                        order = len(elements)

                        element = ParsedElement(
                            element_id=f"elem-{order}",
                            element_type=ElementType.PARAGRAPH,
                            content=text,
                            order=order,
                            heading_level=None,
                            parent_id=parent_id,
                            section_path=section_path,
                            metadata={
                                "style_name": style_name,
                                "character_count": len(text),
                            },
                        )
                        elements.append(element)

                elif isinstance(block, Table):
                    table_content = self._extract_table_content(block)
                    if not table_content:
                        continue

                    parent_id = heading_stack[-1][1] if heading_stack else None
                    section_path = tuple(item[2] for item in heading_stack)
                    order = len(elements)

                    element = ParsedElement(
                        element_id=f"elem-{order}",
                        element_type=ElementType.TABLE,
                        content=table_content,
                        order=order,
                        heading_level=None,
                        parent_id=parent_id,
                        section_path=section_path,
                        metadata={
                            "row_count": len(block.rows),
                            "col_count": len(block.columns),
                            "character_count": len(table_content),
                        },
                    )
                    elements.append(element)

            parser_metadata: dict[str, Any] = {
                "parser_class": self.__class__.__name__,
                "total_elements": len(elements),
                "headings_count": len([e for e in elements if e.element_type == ElementType.HEADING]),
            }

            return ParsedDocument(
                document_id=document.document_id,
                project_id=document.project_id,
                document_type=DocumentType.DOCX,
                elements=elements,
                source_metadata=dict(document.source_metadata),
                parser_metadata=parser_metadata,
                document_version_id=document.document_version_id,
            )
        except Exception as exc:
            if isinstance(exc, (InvalidParsingInputError, DocumentExtractionError)):
                raise
            raise DocumentExtractionError(
                f"Failed to parse DOCX document '{document.document_id}': {str(exc)}",
                original_error=exc,
            ) from exc
